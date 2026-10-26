# -*- coding: utf-8 -*-
# Régénère Doc/mqtt.docx à partir de la lecture des sources (Capteur, Relais-Lora, Home-Assist,
# CydMonitor, Home-Assist-App, configuration.yaml / automations.yaml). Styles repris du document existant.
import sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

DOC = r"D:\Developpement\Pilotage_Maison\Doc\mqtt.docx"
SRC = sys.argv[1] if len(sys.argv) > 1 else DOC  # gabarit de styles
DST = sys.argv[2] if len(sys.argv) > 2 else DOC

doc = Document(SRC)
body = doc.element.body
for el in list(body):
    if el.tag != qn('w:sectPr'):
        body.remove(el)

MONO = 'Consolas'

def shade(cell, color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), color)
    tcPr.append(shd)

def add_runs(p, text, size=None):
    """Texte avec `code` entre accents graves -> police à chasse fixe, **gras**."""
    import re
    for part in re.split(r'(`[^`]*`|\*\*[^*]*\*\*)', text):
        if not part:
            continue
        if part.startswith('`'):
            r = p.add_run(part[1:-1]); r.font.name = MONO
            r._element.rPr.rFonts.set(qn('w:eastAsia'), MONO)
            r.font.size = Pt(9)
        elif part.startswith('**'):
            r = p.add_run(part[2:-2]); r.bold = True
            if size: r.font.size = Pt(size)
        else:
            r = p.add_run(part)
            if size: r.font.size = Pt(size)

def H(text, level=1):
    doc.add_heading(text, level=level)

def P(text, style=None):
    p = doc.add_paragraph(style=style)
    add_runs(p, text)
    return p

def B(text):
    return P(text, 'List Bullet')

def CODE(*lines):
    for l in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Pt(14)
        r = p.add_run(l); r.font.name = MONO; r.font.size = Pt(9)
        r._element.rPr.rFonts.set(qn('w:eastAsia'), MONO)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

def T(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ''
        r = c.paragraphs[0].add_run(h); r.bold = True; r.font.size = Pt(9)
        shade(c, 'D9E2F3')
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''
            add_runs(cells[i].paragraphs[0], v, size=9)
    if widths:
        from docx.shared import Cm
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t

def ENCADRE(titre, lignes, couleur='FFF2CC'):
    t = doc.add_table(rows=1, cols=1)
    t.style = 'Table Grid'
    c = t.rows[0].cells[0]; shade(c, couleur)
    p = c.paragraphs[0]
    r = p.add_run(titre); r.bold = True; r.font.size = Pt(11)
    for l in lines_or(lignes):
        q = c.add_paragraph(); add_runs(q, l, size=10)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

def lines_or(x):
    return x if isinstance(x, (list, tuple)) else [x]

# ════════════════════════════════════════════════════════════════════════════
doc.add_heading('Trames MQTT et LoRa — Pilotage Maison', 0)
P('Inventaire des trames échangées en MQTT et en LoRa P2P entre les capteurs (firmware Capteur), '
  'le Relais-Lora, les CYD (Home-Assist), CydMonitor, l\'application Android et Home Assistant : '
  'syntaxe exacte, topics, et **liste exhaustive des commandes traitées** par chaque récepteur.')
P('Mise à jour du 26/10/2026, établie par lecture des sources : Capteur/src, Relais-Lora/src, Home-Assist/src, '
  'CydMonitor/src, Home-Assist-App 1.2/lib, Doc/configuration.yaml et Doc/automations.yaml.')

ENCADRE('⚠ Commandes destinées à un capteur endormi : retain = true, puis payload vide', [
    'Un capteur en deep sleep n\'est à l\'écoute que pendant ses quelques secondes d\'éveil. '
    'Une commande publiée pendant son sommeil est perdue, sauf si elle est **retenue par le broker**.',
    '1. Publier la commande avec **retain = true** sur le topic de configuration du capteur '
    '(ex. `home/confthremise/command`, payload `SLEEP SLEEPL`). '
    'Le capteur Wi-Fi la reçoit dès son abonnement au réveil (QoS 1) ; il garde ensuite 1 s d\'écoute '
    'après la remontée de ses mesures (DELAI_COMMANDES_APRES_REMONTEE_MS) pour la traiter.',
    '2. **S\'assurer de l\'exécution** : le capteur publie sur `home/<conf>/state` l\'accusé '
    '`<nom> COMMANDE reçue à <ms> ms : <commande> (deep sleep : ACTIF|INACTIF)`, '
    'et au coucher `<nom> DEEPSLEEP ON … <durée>` montre la nouvelle durée.',
    '3. **Effacer la commande retenue** en publiant un **payload vide avec retain = true** sur le même topic. '
    'Sinon elle est rejouée à chaque réveil (et à chaque reconnexion du Relais-Lora). '
    'Un payload vide est ignoré par le capteur (pas d\'accusé) et par le relais.',
    'Précautions :',
    '• Ne jamais retenir sur `home/thermometre/command` : ce topic est partagé par tous les capteurs, '
    'le broker ne garde qu\'UN message retenu par topic (le dernier écrase les autres), et un '
    '`FLOTTEURREMISE ACK` retenu acquitterait d\'avance toutes les remontées suivantes.',
    '• En LoRa, le retain ne suffit pas aujourd\'hui : le Relais-Lora reçoit le message retenu à sa '
    'propre (re)connexion et le réémet aussitôt par radio ; si le capteur dort, il est perdu. '
    'Évolution prévue : le relais mémorise la dernière commande par capteur et la réémet juste après '
    'avoir reçu une trame de ce capteur.',
    '• L\'automatisation HA « Cuve - SLEEP et ACK flotteur » publie actuellement avec retain: false, '
    'en réponse à une trame FlotteurRemise (donc capteur éveillé).',
])

# ────────────────────────────────────────────────────────────────────────────
H('1. Conventions générales')
H('1.1 Broker', 2)
T(['Paramètre', 'Valeur'], [
    ('Serveur / port', '51.77.244.19 / 1883 (identifiants : voir secret.h, non commité)'),
    ('Préfixe des topics', '`home/` (NVS « domo_pref » / « dom_pre »)'),
    ('Topics de configuration', '`home/<sous-topic>command` et `home/<sous-topic>state`, sous-topic en NVS : '
     '`configuration/` (CYD principal), `configaux/`, `configaux2/` (CYD auxiliaires), '
     '`confthremise/`, `confthnomade/` (capteurs), `confrelaislora2/` (relais)'),
    ('QoS', 'Capteur : abonnements QoS 1. CYD, relais : QoS 0. Application : QoS 1 (atLeastOnce).'),
    ('Retain', 'Aucune publication des firmwares n\'est retenue (retain = false partout). '
     'Le retain n\'est utilisé que par l\'émetteur des commandes (voir encadré).'),
])

H('1.2 Format des trames de mesure', 2)
P('Tous les capteurs publient sur le topic commun `home/thermometre/state`. Le premier mot identifie l\'expéditeur '
  '(insensible à la casse côté récepteurs) :')
CODE('<NOM> <MESURE> [<ETAT>] <JJ/MM/AAAA> <HH:MM:SS> <VALEUR(S)> [FORCE] <IP>')
B('`FORCE` : présent quand la mesure a été forcée (commande MESURE, relance, premier envoi après réveil).')
B('`<IP>` : ajoutée en fin de trame par CMqtt::publish() du capteur Wi-Fi ; `0.0.0.0` quand la trame est '
  'arrivée par LoRa (ajoutée par le Relais-Lora si le dernier champ n\'est pas une IP).')
B('Date / heure : `DATE` et `TIME` sont des jetons que le Relais-Lora remplace par sa propre date / heure. '
  'Un capteur LoRa sans heure valide envoie actuellement `Erreur Erreur` (non substitué) — voir § 9.')
B('Suffixe `R` (TEMPR, HUMR, TENSIONR, ONOFFR, COULR, LUXR, DELR, INACTIFR) : '
  'état republié par le CYD principal pour synchroniser les CYD auxiliaires et l\'application.')

# ────────────────────────────────────────────────────────────────────────────
H('2. Trames LoRa P2P')
H('2.1 Paramètres radio', 2)
T(['Paramètre', 'Valeur (Capteur et Relais-Lora)'], [
    ('Module', 'SX1262 (Wio-SX1262 sur XIAO ESP32-S3), RadioLib'),
    ('Fréquence', '915,0 MHz'),
    ('Bande passante / SF / CR', '125 kHz / SF7 / 4/5'),
    ('Sync word / préambule', '0x12 / 8 symboles'),
    ('Puissance', '14 dBm'),
])
H('2.2 Encadrement', 2)
P('Une trame LoRa transporte un couple topic + payload MQTT, encadré par dix signes « = » :')
CODE('==========<topic> <payload>==========')
B('Le caractère `=` est interdit dans le topic et le payload (trame rejetée, code -1).')
B('Le premier mot est le topic, le reste (au moins un champ) est le payload, recomposé avec un espace '
  'entre les champs (les espaces multiples sont réduits à un seul).')
B('Capteur → relais : si le capteur a aussi une IP (mode mixte Wi-Fi + LoRa), il ajoute son IP en fin de payload ; '
  'sinon le relais ajoute `0.0.0.0` avant de publier sur MQTT.')
B('Relais → capteur : le payload est transmis tel que reçu en MQTT (sans IP).')

H('2.3 Trames émises par les capteurs LoRa (relayées sur MQTT)', 2)
T(['Trame LoRa (exemple ThRemise)', 'Origine'], [
    ('`==========home/thermometre/state ThRemise Temp <date> <heure> 21.6 [FORCE]==========`', 'DHT20 — température'),
    ('`==========home/thermometre/state ThRemise Hum <date> <heure> 54.6 [FORCE]==========`', 'DHT20 — humidité'),
    ('`==========home/thermometre/state BatRemise Tension CHARGEE|DECHARGEE <date> <heure> 4.9 [FORCE]==========`', 'Batterie'),
    ('`==========home/thermometre/state FlotteurRemise OnOff <date> <heure> 0|1 [FORCE]==========`', 'Flotteur'),
    ('`==========home/thermometre/state <nom> ALIVE <date> <heure>==========` ou `… ALIVE UPTIME <s>`', 'Watchdog ALIVE (hors deep sleep)'),
    ('`==========home/thermometre/state ThRemise DEEPSLEEP ON DATE HEURE <s> Durée éveil : <ms> ms==========`', 'Mise en sommeil (LoRa seul)'),
    ('`==========home/confthremise/state === ÉTAT <nom> ===\\n<date> <heure>\\n[IP : …\\n]Wdog : …\\nUptime : <s>s==========`', 'Réponse à STATUS'),
])
H('2.4 Trames émises par le Relais-Lora vers les capteurs', 2)
T(['Trame LoRa', 'Condition'], [
    ('`==========home/thermometre/command <payload>==========`',
     'Payload MQTT reçu sur `home/thermometre/command` commençant par THREMISE, BATREMISE, FLOTTEURREMISE, '
     'THNOMADE, BATNOMADE, FLOTTEURNOMADE, THCAVE ou BATCAVE'),
    ('`==========home/confthremise/command <payload>==========`', 'Tout payload reçu sur `home/confthremise/command`'),
])
P('Les commandes que le capteur traite à réception sont celles du § 3.2, identiques en MQTT et en LoRa '
  '(sauf mention contraire).')

# ────────────────────────────────────────────────────────────────────────────
H('3. Capteurs (firmware Capteur)')
P('Noms d\'équipement en service : ThSdb / BatSdb / CoulSdb, ThCh1er / BatCh1er, ThCave / BatCave, '
  'ThChRdc / BatChRdc, ThRemise / BatRemise / FlotteurRemise, ThNomade / BatNomade / FlotteurNomade, '
  'NoeudChauffageSb. Chaque sous-capteur a son nom et son sous-topic en NVS (page web).')
H('3.1 Publications', 2)
T(['Topic', 'Payload', 'Quand'], [
    ('`home/thermometre/state`', '`<nom> Temp <date> <heure> <t°C 0.1> [FORCE] <IP>`', 'DS18B20 / DHT20'),
    ('`home/thermometre/state`', '`<nom> Hum <date> <heure> <%HR 0.1> [FORCE] <IP>`', 'DHT20'),
    ('`home/thermometre/state`', '`<nom> Tension CHARGEE|DECHARGEE <date> <heure> <V 0.1> [FORCE] <IP>`', 'Batterie AA (DECHARGEE sous la tension mini)'),
    ('`home/thermometre/state`', '`<nom> OnOff <date> <heure> 0|1 [FORCE] <IP>`', 'Flotteur (TOR)'),
    ('`home/thermometre/state`', '`<nom> Coul <date> <heure> <R> <G> <B> <Lux> <LuxBrut> [FORCE] <IP>`', 'TCS34725'),
    ('`home/thermometre/state`', '`<nom> Del <date> <heure> 0|1 [FORCE] <IP>`', 'TCS34725 (détection DEL)'),
    ('`home/noeud/state`', '`<nom> OnOff <date> <heure> 0|1 [FORCE] <IP>`', 'Micro MAX4466 (état chauffage)'),
    ('`home/noeud/state`', '`<nom> IRKO <date> <heure> <IP>`', 'Échec de l\'actionneur IR'),
    ('topic state du sous-capteur', '`<nom> ALIVE <date> <heure> [RSSI <dBm>] <IP>` ou `<nom> ALIVE UPTIME <s> …`', 'Watchdog ALIVE (période WDOG, désactivé en deep sleep)'),
    ('`home/thermometre/state`', '`<nomConfig> REVEIL DEEPSLEEP <date> <heure> <s dormies> Cause réveil : … Cause reset : … [RSSI : … dBm] [Préc : …] <IP>`', 'Réveil après deep sleep (Wi-Fi)'),
    ('`home/thermometre/state`', '`<nomConfig> REVEIL OFF <date> <heure> Cause réveil : … <IP>`', 'Démarrage sans deep sleep'),
    ('`home/thermometre/state`', '`<nomConfig> DEEPSLEEP ON <date> <heure> <s> Durée éveil : <ms> ms <IP>`', 'Mise en sommeil'),
    ('`home/thermometre/state`', '`<nomConfig> online le <date> à <heure> <IP>`', 'Connexion MQTT (Wi-Fi)'),
    ('`home/<conf>/state`', '`<nomConfig> offline`', 'Testament (LWT), non retenu'),
    ('`home/<conf>/state`', '`<nomConfig> COMMANDE reçue à <ms> ms : <commande> (deep sleep : ACTIF|INACTIF)`', 'Accusé de toute commande de configuration non vide (MQTT seulement)'),
    ('`home/<conf>/state`', 'Rapport multiligne `=== ÉTAT <nom> ===` (date, IP, Wdog, Uptime)', 'Réponse à STATUS'),
    ('`home/<conf>/state/echo`', '`<nomConfig>`', 'Écho interne avant coupure (le capteur s\'y abonne lui-même)'),
    ('`home/configuration/command`', '`NONCONFIGURE <modèle>-<fin MAC> [<date> <heure>] <IP>`', 'Puce neuve non configurée (toutes les 2 s puis 60 s)'),
])

H('3.2 Commandes traitées par un capteur (liste exhaustive)', 2)
P('Casse indifférente. Le premier mot doit être exactement le nom de l\'équipement visé. '
  'Une commande non listée est ignorée (trace série « Commande inconnue » ou « Non traité »).')
T(['Topic', 'Commande', 'Effet', 'MQTT / LoRa'], [
    ('`home/<conf>/command` (ex. confthremise)', '`SLEEP ENABLE`', 'Active le deep sleep (NVS)', 'les deux'),
    ('', '`SLEEP DISABLE` ou `SLEEP -1`', 'Désactive le deep sleep', 'les deux'),
    ('', '`SLEEP <n>`', 'Intervalle de sommeil = n s, et active le deep sleep', 'les deux'),
    ('', '`SLEEP SLEEPS` / `SLEEP SLEEPM` / `SLEEP SLEEPL`', 'Applique la durée Small / Medium / Large, et active le deep sleep', 'les deux'),
    ('', '`SLEEP SLEEPS <n>` (idem M, L)', 'Règle la durée Small (Medium, Large) à n s, sans l\'appliquer', 'les deux'),
    ('', '`WDOG ENABLE` / `WDOG DISABLE`', 'Active / désactive le watchdog ALIVE', 'les deux'),
    ('', '`WDOG <s>`', 'Période ALIVE = s secondes (> 0), et active le watchdog', 'les deux'),
    ('', '`STATUS`', 'Publie le rapport d\'état sur `home/<conf>/state` (par LoRa si reçu par LoRa)', 'les deux'),
    ('', '`REBOOT`', 'Redémarre l\'ESP32', 'les deux'),
    ('`home/thermometre/command`', '`<nomCapteur> MESURE`', 'Mesure immédiate et publication FORCE (DS18B20, DHT20, TCS34725, flotteur)', 'les deux'),
    ('', '`<nomFlotteur> ACK`', 'Acquitte la remontée du flotteur : nécessaire au coucher si l\'option « Acknoledge » est cochée (LoRa)', 'les deux'),
    ('', '`<nomCapteur> SLEEP <n>`', 'Analysée mais sans effet (utiliser le topic de configuration)', 'les deux'),
    ('`home/noeud/command` (sous-topic du micro)', '`<nomMicro> MESURE`', 'Publie l\'état OnOff du chauffage', 'MQTT'),
    ('', '`<nomMicro> ON` / `<nomMicro> OFF`', 'Allumage / extinction du chauffage par l\'actionneur IR (si ACTIONNEUR_IR)', 'MQTT'),
])
P('Attention à la syntaxe : `SLEEPS` seul est refusé, il faut `SLEEP SLEEPS`. Les durées sont en secondes.')

# ────────────────────────────────────────────────────────────────────────────
H('4. Relais-Lora (passerelle MQTT ↔ LoRa)')
H('4.1 Abonnements et commandes traitées (liste exhaustive)', 2)
T(['Topic écouté', 'Commande / filtre', 'Effet'], [
    ('`home/confrelaislora2/command`', '`STATUS`', 'Publie `=== ÉTAT <nom> ===\\n<date> <heure>\\nIP : <ip>\\nUptime : <s>s` sur `home/confrelaislora2/state`'),
    ('', '`REBOOT`', 'Redémarre le relais'),
    ('`home/thermometre/command`', 'Payload commençant par THREMISE, BATREMISE, FLOTTEURREMISE, THNOMADE, BATNOMADE, FLOTTEURNOMADE, THCAVE, BATCAVE', 'Réémis par LoRa (§ 2.4) ; autres payloads ignorés'),
    ('`home/confthremise/command`', 'Tout payload', 'Réémis par LoRa (§ 2.4)'),
    ('LoRa (réception)', 'Toute trame valide `==========<topic> <payload>==========`', 'Publiée sur MQTT `<topic>` avec DATE / TIME substitués et `0.0.0.0` ajouté'),
])
H('4.2 Publications propres', 2)
T(['Topic', 'Payload', 'Quand'], [
    ('`home/configuration/command` (topic du maître)', '`DEVICE RELAIS LORA <nom> <date> <heure>`', 'Démarrage'),
    ('`home/confrelaislora2/state`', '`<nom> - Reconnect le <date> à <heure>`', 'Testament (LWT)'),
    ('`home/confrelaislora2/state`', '`<nom> REVEIL OFF <date> <heure>` / `<nom> REVEIL DEEPSLEEP …`', 'Démarrage'),
    ('`home/confrelaislora2/state`', 'Rapport STATUS', 'Réponse à STATUS'),
])

# ────────────────────────────────────────────────────────────────────────────
H('5. CYD Home-Assist (principal et auxiliaires)')
P('Le CYD principal est en __LOCAL_MODE__ (pilote chaudière et équipements 433 MHz) ; les CYD auxiliaires '
  'sont en mode REMOTE et suivent les états republiés par le principal.')
H('5.1 Commandes traitées sur le topic de configuration (liste exhaustive)', 2)
P('Topic `home/configuration/command` pour le principal (`home/configaux/command`, `home/configaux2/command` pour les auxiliaires).')
T(['Commande', 'Effet', 'Réponse'], [
    ('`STATUS`', 'Rapport : état chaudière / projecteur / guirlande / SDB, veille, IP, uptime, contrôleurs connus', 'Multiligne sur `…/state`'),
    ('`VEILLE <s>`', 'Délai de mise en veille de l\'écran (s)', '`VEILLE <s>` sur `…/state`'),
    ('`ECRAN SLEEP` / `ECRAN WAKE`', 'Éteint / rallume l\'écran', '—'),
    ('`ECRAN SERIE <n>`', 'Série d\'écrans affichée', '`ECRAN SERIE <n>` sur `…/state`'),
    ('`ECRAN <n>`', 'Orientation 0 à 3 (portrait, paysage, inversés)', '`ECRAN <n>` sur `…/state`'),
    ('`ALIVE <min>`', 'Période du watchdog ALIVEWDOG (minutes)', 'message d\'état à l\'écran'),
    ('`DEVICE CYD <nom> <date> <heure> 0 <IP>`', 'Principal : note le CYD auxiliaire et republie tous les états', 'États ONR/OFFR/ENABLER/DISABLER, TEMPR, HUMR…'),
    ('`DEVICE APP <téléphone> <date heure> 0 <IP>`', 'Principal : affiche l\'application (série 5) et republie tous les états', 'idem'),
    ('`DEVICE RELAIS LORA <nom> <date> <heure>`', 'Principal : affiche le relais et republie tous les états', 'idem'),
    ('`NONCONFIGURE <modèle>-<MAC> [<date> <heure>] <IP>`', 'Affiche la puce neuve (série 5 et barre d\'état)', '—'),
    ('`REBOOT`', 'Redémarre le CYD', '—'),
])
P('Le principal écoute aussi `home/configuration/state` : `<nom du CYD> ALIVEWDOG <date> <heure> <IP>` '
  '(note le contact d\'un CYD auxiliaire ; le nom peut contenir des espaces).')

H('5.2 Commandes traitées par les équipements pilotés (liste exhaustive)', 2)
T(['Topic', 'Commande reçue', 'Effet', 'Réponse (même topic)'], [
    ('`home/chaudiere/command`', '`ON` / `OFF`', 'Marche / arrêt (si active)', '`ONR` / `OFFR`'),
    ('', '`ON_FORCE` / `OFF_FORCE`', 'Force l\'activation (ENABLE) puis ON / force la désactivation', '`ENABLER` / `DISABLER` + `ONR` / `OFFR`'),
    ('', '`ENABLE` / `DISABLE`', 'Active / désactive l\'équipement', '`ENABLER` / `DISABLER`'),
    ('`home/projecteur/command`, `home/guirlande/command`', '`ON` / `OFF` / `ENABLE` / `DISABLE`', 'Émission RC 433 MHz', '`ONR` / `OFFR` / `ENABLER` / `DISABLER`'),
    ('`home/chauffage/command`', '`ON` / `OFF` / `ON_FORCE` / `OFF_FORCE` / `ENABLE` / `DISABLE`', 'Chauffage SDB : relais + séquence IR via le noeud (publie `NoeudChauffageSb ON|OFF` sur `home/noeud/command`)', '`ONR` / `OFFR` / `ENABLER` / `DISABLER`'),
    ('mêmes topics (CYD auxiliaire)', '`ONR` / `OFFR` / `ENABLER` / `DISABLER`', 'Mise à jour des boutons', '—'),
])
P('Particularité : les réponses d\'état sont publiées sur le topic **command** (pas sur state). '
  'Home Assistant (configuration.yaml) et l\'application utilisent donc `…/command` comme state_topic.')

H('5.3 États reçus sur home/thermometre/state', 2)
P('Identifiants reconnus (début de payload) : THCYD, THSDB, COULSDB, BATSDB, THCH1ER, BATCH1ER, THCAVE, BATCAVE, '
  'THREMISE, BATREMISE, FLOTTEURREMISE, THNOMADE, BATNOMADE, FLOTTEURNOMADE, THCHRDC, BATCHRDC. Autres : ignorés.')
T(['Mesure (2e mot)', 'Équipements', 'Mode'], [
    ('`Temp`, `Hum`', 'thermomètres', 'tous'),
    ('`Tension CHARGEE|DECHARGEE`', 'batteries', 'tous'),
    ('`OnOff`', 'flotteurs', 'tous'),
    ('`Coul`, `Lux`, `Del`', 'COULSDB', 'tous'),
    ('`ALIVE`', 'tous (réinitialise le watchdog)', 'tous'),
    ('`TEMPR`, `HUMR`, `TENSIONR`, `ONOFFR`, `COULR`, `LUXR`, `DELR`, `INACTIFR`', 'états republiés par le principal', 'auxiliaires'),
])
H('5.4 Autres topics écoutés', 2)
T(['Topic', 'Payload', 'Effet'], [
    ('`home/noeud/state`', '`NOEUDCHAUFFAGESB OnOff|Etat <date> <heure> 0|1 <IP>`', 'État du chauffage SDB'),
    ('', '`NOEUDCHAUFFAGESB ALIVE …` / `NOEUDCHAUFFAGESB IRKO …`', 'Watchdog / échec IR (coupure du relais)'),
    ('', '`NOEUDCHAUFFAGESB ONOFFR|INACTIFR …`', 'Auxiliaires seulement'),
    ('`home/homeassistant/state`', '`NewNas ETAT <date> <heure> 0|1` / `BigNas ETAT …`', 'État des NAS (publié par Home Assistant)'),
])
H('5.5 Publications du CYD', 2)
T(['Topic', 'Payload', 'Quand'], [
    ('`home/thermometre/command`', '`<nom> MESURE`', 'Demande de mesure (topic de commande de l\'objet distant, NVS)'),
    ('`home/thermometre/state`', '`<nom> TEMPR|HUMR|TENSIONR|ONOFFR|COULR|LUXR <date> <heure> <val>` / `<nom> INACTIFR <date> <heure>`', 'Principal : synchronisation des auxiliaires'),
    ('`home/thermometre/state`', '`<nom> Temp <date> <heure> <t> [FORCE]`', 'Capteur local DS18B20 / DHT20 du CYD'),
    ('`home/configuration/state` (topic du maître)', '`<nom> ALIVEWDOG <date> <heure> <IP>`', 'Watchdog périodique (auxiliaire : sur le topic d\'état du maître)'),
    ('`home/configuration/command` (topic du maître)', '`DEVICE CYD <nom> <date> <heure> 0 <IP>`', 'Démarrage d\'un CYD auxiliaire'),
    ('`home/<conf>/state`', '`VEILLE <s>`, `ECRAN SERIE <n>`, `ECRAN <n>`, rapport STATUS', 'Réponses / démarrage'),
    ('`home/<conf>/state`', '`<nom> - Reconnect le <date> à <heure>`', 'Testament (LWT)'),
    ('`home/log`', 'texte libre', 'Journal de débogage (MyLogger)'),
])

# ────────────────────────────────────────────────────────────────────────────
H('6. Autres clients MQTT')
H('6.1 Application Android (Home-Assist-App)', 2)
B('Publie `DEVICE APP <téléphone> <date heure> 0 <IP>` sur `home/configuration/command` au démarrage.')
B('Publie `ON` / `OFF` / `ENABLE` / `DISABLE` sur `home/<équipement>/command` (chaudiere, projecteur, guirlande, chauffage).')
B('Écoute `home/thermometre/state` (Temp, Hum, Tension, OnOff, Coul, Del et variantes R ; les BATxxx sont rattachés au THxxx) '
  'et `home/<équipement>/command` (ONR, OFFR, ENABLER, DISABLER).')
H('6.2 CydMonitor', 2)
B('Écoute uniquement `home/thermometre/state`, ne publie rien.')
B('Ne traite que les trames dont le premier mot est exactement `ThNomade` : `ThNomade Del … <0|1>` (disque vert / gris) '
  'et `ThNomade Coul … <R> <G> <B> <Lux> <Brut>`.')
H('6.3 Home Assistant (configuration.yaml, automations.yaml)', 2)
T(['Automatisation / entité', 'Trame'], [
    ('Interrupteurs', '`ON`/`OFF`, `ENABLE`/`DISABLE`, `ON_FORCE`/`OFF_FORCE` sur `home/{chaudiere,projecteur,guirlande,chauffage}/command`, état lu sur le même topic'),
    ('Capteurs', 'Gabarits sur `home/thermometre/state` ; IP = dernier mot du payload'),
    ('CYD - Changer veille / série / orientation / intervalle Alive', '`VEILLE <s>`, `ECRAN SERIE <n>`, `ECRAN <n>`, `ALIVE <min>` sur `home/configuration|configaux|configaux2/command`'),
    ('CYD - Recevoir …', 'Lit `VEILLE`, `ECRAN SERIE`, `ECRAN`, `ALIVEWDOG`, `ALIVE` et STATUS sur les topics state'),
    ('NAS - Publier état MQTT', '`NewNas|BigNas ETAT <date> <heure> 0|1` sur `home/homeassistant/state`'),
    ('HA - Publier IP', '`PIOS IO <date> <heure> <IP>` sur `home/homeassistant/state`'),
    ('Cuve - SLEEP et ACK flotteur', 'Sur trame FLOTTEURREMISE : `SLEEP SLEEPS` ou `SLEEP SLEEPL` (selon la pluie prévue) sur `home/confthremise/command`, puis `FLOTTEURREMISE ACK` sur `home/thermometre/command` (retain: false)'),
    ('Nomade - SLEEP et ACK flotteur', 'Idem avec `home/confthnomade/command` et `FLOTTEURNOMADE ACK`'),
    ('Chaudière (automatisations)', '`ON_FORCE` / `OFF_FORCE` sur `home/chaudiere/command`'),
])

# ────────────────────────────────────────────────────────────────────────────
H('7. Récapitulatif des topics')
T(['Topic', 'Publié par', 'Écouté par'], [
    ('`home/thermometre/state`', 'Capteurs (Wi-Fi, ou LoRa via relais), CYD principal (R, local)', 'CYD, application, CydMonitor, Home Assistant'),
    ('`home/thermometre/command`', 'CYD (MESURE), Home Assistant (ACK)', 'Capteurs Wi-Fi, Relais-Lora (→ LoRa)'),
    ('`home/noeud/state`', 'Noeud chauffage (micro + IR)', 'CYD'),
    ('`home/noeud/command`', 'CYD (ON/OFF, MESURE)', 'Noeud chauffage'),
    ('`home/{chaudiere,projecteur,guirlande,chauffage}/command`', 'HA, application, CYD auxiliaires (commandes) ; CYD principal (réponses R)', 'CYD, HA, application'),
    ('`home/configuration/command`', 'HA, application (DEVICE APP), CYD aux (DEVICE CYD), relais (DEVICE RELAIS LORA), puces neuves (NONCONFIGURE)', 'CYD principal'),
    ('`home/configuration/state`', 'CYD principal et auxiliaires (ALIVEWDOG, réponses)', 'CYD principal, HA'),
    ('`home/configaux/…`, `home/configaux2/…`', 'HA / CYD auxiliaires', 'CYD auxiliaires / HA'),
    ('`home/confthremise/command`', 'HA (SLEEP)', 'ThRemise (Wi-Fi), Relais-Lora (→ LoRa)'),
    ('`home/confthnomade/command`', 'HA (SLEEP)', 'ThNomade (Wi-Fi) — non relayé en LoRa'),
    ('`home/conf<capteur>/state`', 'Capteurs (accusé, STATUS, LWT)', 'Home Assistant, MQTTBox'),
    ('`home/confrelaislora2/command|state`', 'HA / relais', 'relais / HA'),
    ('`home/homeassistant/state`', 'Home Assistant', 'CYD'),
    ('`home/log`', 'CYD', 'débogage'),
])

# ────────────────────────────────────────────────────────────────────────────
H('8. Ce qui a changé par rapport à la version précédente du document')
B('Les réponses d\'état des équipements (ONR, OFFR, ENABLER, DISABLER) sont sur le topic **command**, pas state.')
B('Le chauffage SDB est sur `home/chauffage/command` (et non chauffagesb).')
B('REVEIL / DEEPSLEEP des capteurs : publiés sur `home/thermometre/state` (topic de leurs sous-capteurs), pas sur le topic de configuration.')
B('Ajouts : trames LoRa, commandes SLEEP / WDOG / STATUS / REBOOT des capteurs, ACK du flotteur, accusé COMMANDE reçue, '
  'ALIVE, NONCONFIGURE, DEVICE CYD / APP / RELAIS LORA, ALIVEWDOG, noeud chauffage, NAS, variantes R.')

# ────────────────────────────────────────────────────────────────────────────
H('9. Anomalies relevées à la lecture des sources')
B('Relais-Lora : pas d\'abonnement à `home/confthnomade/command` ; les SLEEP envoyés à ThNomade par HA ne sont pas relayés en LoRa.')
B('Relais-Lora : aucune file d\'attente ; une commande arrivant pendant le sommeil du capteur LoRa est perdue (voir encadré).')
B('Capteur LoRa : `sendPacket()` remet `operationDone` à false avant d\'émettre, ce qui peut effacer une réception pas encore traitée.')
B('Capteur LoRa sans heure : les mesures partent avec `Erreur Erreur` au lieu des jetons `DATE TIME`, et DEEPSLEEP ON avec `DATE HEURE` ; '
  'le relais ne substitue que DATE et TIME.')
B('Batterie : aucune commande MESURE n\'est traitée côté capteur (pas d\'aiguillage BATxxx sur `home/thermometre/command`), '
  'alors que le CYD en envoie.')
B('CConfig::handleMqttCommand() du Relais-Lora est vide : seuls STATUS et REBOOT agissent.')

doc.save(DST)
print('OK', DST)
